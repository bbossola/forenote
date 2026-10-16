"""DOCX, PPTX, and XLSX."""

from __future__ import annotations

from pathlib import Path

from forenote.kb.parsers import Limits, ParseFailure
from forenote.kb.parsers.sheets import row_text, sheet_sections
from forenote.kb.sections import Section

_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _check_container(path: Path) -> None:
    """An encrypted OOXML file is an OLE container, not a ZIP file."""
    with path.open("rb") as f:
        if f.read(8) == _OLE:
            raise ParseFailure("encrypted")


def parse_docx(path: Path, limits: Limits) -> list[Section]:
    _check_container(path)
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []

    def flush() -> None:
        text = "\n".join(buf).strip()
        if text:
            sections.append(Section(text, {"heading": heading} if heading else {}))
        buf.clear()

    for block in document.iter_inner_content():
        if hasattr(block, "rows"):  # a table
            for row in block.rows:
                buf.append("\t".join(cell.text for cell in row.cells))
            continue
        style = block.style.name if block.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            flush()
            heading = block.text.strip() or heading
        buf.append(block.text)
    flush()
    return sections


def _shape_texts(shape):
    from pptx.shapes.group import GroupShape

    if isinstance(shape, GroupShape):
        for inner in shape.shapes:
            yield from _shape_texts(inner)
        return
    if shape.has_text_frame:
        yield shape.text_frame.text
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            yield "\t".join(cell.text for cell in row.cells)


def parse_pptx(path: Path, limits: Limits) -> list[Section]:
    _check_container(path)
    import pptx

    slides = list(pptx.Presentation(str(path)).slides)
    if len(slides) > limits.max_pages:
        raise ParseFailure("too_many_pages")
    out = []
    for number, slide in enumerate(slides, 1):
        parts = [t for shape in slide.shapes for t in _shape_texts(shape)]
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                parts.append(frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text:
            out.append(Section(text, {"slide": number}))
    return out


def parse_xlsx(path: Path, limits: Limits) -> list[Section]:
    _check_container(path)
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        out: list[Section] = []
        for ws in wb.worksheets:
            rows = ((n, row_text(values)) for n, values in enumerate(ws.iter_rows(values_only=True), 1))
            out += sheet_sections(ws.title, rows)
        return out
    finally:
        wb.close()
