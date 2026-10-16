"""Build the format fixtures of the knowledge base tests. We wrote all the text.

Run once and commit the output:
    uv run python tests/fixtures/kb/make_fixtures.py
"""

from __future__ import annotations

import io
from pathlib import Path

HERE = Path(__file__).resolve().parent

PDF_PAGES = [
    "Forenote sample PDF.\nThe budget for the third quarter is 1200 euros.",
    "",
    "Page three is about hiring.\nTwo engineers join in May.",
]


def make_pdf(pages: list[str]) -> bytes:
    """A small valid PDF with one Helvetica text block on each page."""
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pages_id = font + 2 * len(pages) + 1
    page_ids = []
    for text in pages:
        ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        for line in text.split("\n"):
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({esc}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (pages_id, content, font)
        ))
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    assert add(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_ids)) == pages_id
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)
    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(buf.tell())
        buf.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = buf.tell()
    buf.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        buf.write(b"%010d 00000 n \n" % off)
    buf.write(b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, catalog, xref))
    return buf.getvalue()


def pdfs() -> None:
    import pypdf

    (HERE / "sample.pdf").write_bytes(make_pdf(PDF_PAGES))
    writer = pypdf.PdfWriter(clone_from=str(HERE / "sample.pdf"))
    writer.encrypt("secret", algorithm="RC4-128")
    writer.write(HERE / "encrypted.pdf")


def docx_file() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Budget", 1)
    d.add_paragraph("The budget for the third quarter is 1200 euros.")
    d.add_heading("Hiring", 2)
    d.add_paragraph("Two engineers join in May.")
    table = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Item", "Cost"], ["Laptop", "1200"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.save(HERE / "sample.docx")
    # An encrypted OOXML file is an OLE container. The parser checks the first 8 bytes.
    (HERE / "encrypted.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504))


def pptx_file() -> None:
    import pptx

    pres = pptx.Presentation()
    one = pres.slides.add_slide(pres.slide_layouts[1])
    one.shapes.title.text = "Roadmap"
    one.placeholders[1].text = "Ship the beta in May"
    one.notes_slide.notes_text_frame.text = "Speaker note about risks"
    two = pres.slides.add_slide(pres.slide_layouts[5])
    two.shapes.title.text = "Thank you"
    pres.save(HERE / "sample.pptx")


def xlsx_file() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    q3 = wb.active
    q3.title = "Q3"
    for row in (["Item", "Cost"], ["Laptop", 1200], [None, None], ["Desk", 300.5]):
        q3.append(row)
    notes = wb.create_sheet("Notes")
    notes.append(["All costs in euros"])
    wb.save(HERE / "sample.xlsx")


def odf_files() -> None:
    from odf import draw, style, table
    from odf import text as odtext
    from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText

    doc = OpenDocumentText()
    doc.text.addElement(odtext.H(outlinelevel=1, text="Budget"))
    doc.text.addElement(odtext.P(text="The budget grows."))
    doc.text.addElement(odtext.H(outlinelevel=2, text="Hiring"))
    doc.text.addElement(odtext.P(text="Two engineers join."))
    doc.save(str(HERE / "sample.odt"))

    pres = OpenDocumentPresentation()
    layout = style.PageLayout(name="layout")
    pres.automaticstyles.addElement(layout)
    master = style.MasterPage(name="master", pagelayoutname=layout)
    pres.masterstyles.addElement(master)
    for name, words in (("s1", "Roadmap slide"), ("s2", "Second slide")):
        page = draw.Page(name=name, masterpagename=master)
        frame = draw.Frame(width="10cm", height="2cm", x="1cm", y="1cm")
        box = draw.TextBox()
        box.addElement(odtext.P(text=words))
        frame.addElement(box)
        page.addElement(frame)
        pres.presentation.addElement(page)
    pres.save(str(HERE / "sample.odp"))

    sheet = OpenDocumentSpreadsheet()
    t = table.Table(name="Q3")
    for values in (["Item", "Cost"], ["Laptop", "1200"]):
        tr = table.TableRow()
        for value in values:
            cell = table.TableCell(valuetype="string")
            cell.addElement(odtext.P(text=value))
            tr.addElement(cell)
        t.addElement(tr)
    sheet.spreadsheet.addElement(t)
    sheet.save(str(HERE / "sample.ods"))


def broken_files() -> None:
    for ext in ("pdf", "docx", "pptx", "xlsx", "odt", "odp", "ods"):
        (HERE / f"broken.{ext}").write_bytes(b"this is not a valid file\n")


if __name__ == "__main__":
    pdfs()
    docx_file()
    pptx_file()
    xlsx_file()
    odf_files()
    broken_files()
